"""Build the requested illustrated Word articles from the approved Markdown."""
from __future__ import annotations
import json
import re
from pathlib import Path
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
BUILD = ROOT / "word-build"
FIG = BUILD / "figures"
HISTORY = ROOT.parent / "moleculenet-reproduction-20260918"
PREVIEW = ROOT.parent / "paper2any-direct-20260919/build/final-preview"
BODY_FONT = "Microsoft YaHei"
TEXT_COLOR = "172534"
WIDTH = 7.02
HEADINGS = {
    "一、名字相近，两个项目解决的不是同一件事": "1 两个项目的职责与接入方式",
    "二、为什么选 MoleculeNet，而不是只让模型总结一篇新论文": "2 MoleculeNet 的测试范围与执行流程",
    "三、Paper2Agent 跑通了，但“跑通”的含义需要说准确": "3 Paper2Agent 的结果与审查限制",
    "四、Paper2Any 的失败，比一份漂亮的成品更值得看": "4 Paper2Any 的生成失败与导出成功",
    "五、实验结果没有配合演示稿“变好看”": "5 真实实验结果与数据纠正",
    "六、这次测试能给出的结论": "6 实测结论与适用范围",
    "一、先看调用链：没有必要再装一套 Pi 核心": "1 一套 Pi 核心的调用架构",
    "二、环境分开准备，先满足最小任务": "2 Python 与命令执行环境",
    "三、安装的是完整 Skill 文件夹，不是一份 Markdown": "3 安装完整的 Paper2Agent Skill",
    "四、从 GUI 发起任务：先生成阅读包，再谈实验": "4 在 GUI 中发起论文任务",
    "五、展开底层命令：哪些可以自动跑，哪些必须审查": "5 转换命令与审查步骤",
    "六、如果还要复现实验，再准备科学环境": "6 独立的科学实验环境",
    "七、可视化测试发现的问题，不能藏在教程后面": "7 可视化测试发现的问题",
    "八、交付的不应只有一条“完成了”": "8 交付文件与验收依据",
    "1. 创建来源快照": "创建来源快照",
    "2. 提取并生成审查辅助材料": "提取与审查",
    "3. 审查完成后再构建和验证": "构建与严格验证",
}

def set_font(run, size=None, bold=None, color=TEXT_COLOR, mono=False):
    run.font.name = "Consolas" if mono else BODY_FONT
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for attr in ("eastAsia", "cs"):
        fonts.set(qn(f"w:{attr}"), BODY_FONT)
    for attr in ("ascii", "hAnsi"):
        fonts.set(qn(f"w:{attr}"), "Consolas" if mono else BODY_FONT)
    for attr in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        fonts.attrib.pop(qn(f"w:{attr}"), None)
    lang = OxmlElement("w:lang")
    lang.set(qn("w:val"), "zh-CN")
    lang.set(qn("w:eastAsia"), "zh-CN")
    rpr.append(lang)

def style_font(style, size, bold=False):
    style.font.name = BODY_FONT
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.italic = False
    style.font.color.rgb = RGBColor(0, 0, 0)
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{attr}"), BODY_FONT)
    for attr in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        fonts.attrib.pop(qn(f"w:{attr}"), None)

def hyperlink(paragraph, label, url, size=None):
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True))
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "246C91")
    rpr.append(color)
    fonts = OxmlElement("w:rFonts")
    for attr in ("ascii", "hAnsi", "eastAsia"):
        fonts.set(qn(f"w:{attr}"), BODY_FONT)
    rpr.append(fonts)
    if size:
        sz = OxmlElement("w:sz")
        sz.set(qn("w:val"), str(round(size * 2)))
        rpr.append(sz)
    run.append(rpr)
    text = OxmlElement("w:t")
    text.text = label
    run.append(text)
    link.append(run)
    paragraph._p.append(link)

INLINE = re.compile(r"(\[[^\]]+\]\(https?://[^)]+\)|\*\*[^*]+\*\*|`[^`]+`|\*[^*]+\*)")

def inline(paragraph, content, size=None, color=TEXT_COLOR):
    for part in INLINE.split(content):
        if not part:
            continue
        if part.startswith("[") and re.fullmatch(r"\[[^\]]+\]\(https?://[^)]+\)", part):
            match = re.fullmatch(r"\[([^\]]+)\]\(([^)]+)\)", part)
            hyperlink(paragraph, match[1], match[2], size)
        else:
            bold = part.startswith("**") and part.endswith("**")
            mono = part.startswith("`") and part.endswith("`")
            italic = part.startswith("*") and part.endswith("*") and not bold
            value = part[2:-2] if bold else part[1:-1] if mono or italic else part
            run = paragraph.add_run(value)
            set_font(run, size, bold, color, mono)
            if italic:
                run.italic = True

def setup(title, subtitle):
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin, section.bottom_margin = Inches(.65), Inches(.65)
    section.left_margin, section.right_margin = Inches(.74), Inches(.74)
    section.footer_distance = Inches(.25)
    for name, size, bold in (
        ("Normal", 11.2, False), ("Title", 25, True), ("Subtitle", 12, False),
        ("Heading 1", 16, True), ("Heading 2", 13, True), ("Heading 3", 11.6, True),
        ("Caption", 9.3, False), ("List Bullet", 11.2, False),
    ):
        style_font(doc.styles[name], size, bold)
    normal = doc.styles["Normal"].paragraph_format
    normal.space_after = Pt(7)
    normal.line_spacing = 1.34
    normal.widow_control = True
    word_wrap = OxmlElement("w:wordWrap")
    word_wrap.set(qn("w:val"), "0")
    doc.styles["Normal"].element.get_or_add_pPr().append(word_wrap)
    for name in ("Heading 1", "Heading 2", "Heading 3"):
        pf = doc.styles[name].paragraph_format
        pf.space_before, pf.space_after = Pt(16), Pt(8)
        pf.keep_with_next = True
        pf.keep_together = True
    doc.styles["Title"].paragraph_format.space_after = Pt(12)
    doc.styles["Subtitle"].paragraph_format.space_after = Pt(11)
    doc.styles["List Bullet"].paragraph_format.space_after = Pt(5)
    doc.styles["Caption"].paragraph_format.space_after = Pt(11)
    doc.styles["Caption"].paragraph_format.line_spacing = 1.2
    for name in ("Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3"):
        ppr = doc.styles[name].element.find(qn("w:pPr"))
        if ppr is not None:
            for border in list(ppr.findall(qn("w:pBdr"))):
                ppr.remove(border)
    code = doc.styles.add_style("Article Code", 1)
    style_font(code, 9.5)
    code.font.name = "Consolas"
    cpf = code.paragraph_format
    cpf.left_indent = Inches(.12)
    cpf.right_indent = Inches(.08)
    cpf.space_before = Pt(0)
    cpf.space_after = Pt(1.5)
    cpf.line_spacing = 1.10
    cpf.widow_control = True
    code_wrap = OxmlElement("w:wordWrap")
    code_wrap.set(qn("w:val"), "1")
    code.element.get_or_add_pPr().append(code_wrap)
    title_par = doc.add_paragraph(title, "Title")
    for run in title_par.runs:
        set_font(run, 25, True, "000000")
    subtitle_par = doc.add_paragraph(subtitle, "Subtitle")
    for run in subtitle_par.runs:
        set_font(run, 12, color="000000")
    author = doc.add_paragraph()
    author.paragraph_format.space_after = Pt(15)
    set_font(author.add_run("Stella    2026 年 9 月 19 日"), 10.2, color="000000")
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(footer.add_run("Stella    "), 9, color="5F6F7A")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    doc.core_properties.author = "Stella"
    doc.core_properties.title = title
    doc.core_properties.subject = subtitle
    doc.core_properties.keywords = "Paper2Agent, Paper2Any, Pi GUI, MoleculeNet"
    doc.core_properties.comments = ""
    return doc

class Article:
    def __init__(self, number, title, subtitle):
        self.doc = setup(title, subtitle)
        self.number = number
        self.figure_count = 0
        self.figures = []
        self.chapter = 0
        self.table_index = 0

    def paragraph(self, text, style=None):
        p = self.doc.add_paragraph(style=style)
        inline(p, text)
        if text.rstrip().endswith(("：", ":")):
            p.paragraph_format.keep_with_next = True
        return p

    def image(self, path, caption, width=WIDTH, space_before=7, caption_after=11):
        path = Path(path)
        if not path.is_file():
            raise FileNotFoundError(path)
        self.figure_count += 1
        p = self.doc.add_paragraph()
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        picture = p.add_run().add_picture(str(path), width=Inches(width))
        picture._inline.docPr.set("descr", caption)
        picture._inline.docPr.set("title", f"图 {self.figure_count}")
        cap = self.doc.add_paragraph(style="Caption")
        cap.paragraph_format.space_after = Pt(caption_after)
        inline(cap, f"图 {self.figure_count}  {caption}", size=9.3, color="526579")
        self.figures.append({"number": self.figure_count, "path": str(path), "caption": caption})

    def table(self, lines):
        rows = [[v.strip() for v in line.strip().strip("|").split("|")] for line in lines]
        rows = [row for row in rows if not all(re.fullmatch(r":?-+:?", v) for v in row)]
        cols = len(rows[0])
        if any(len(row) != cols for row in rows):
            raise ValueError("Irregular Markdown table")
        self.table_index += 1
        if cols == 4:
            widths = [2.12, 1.63, 1.63, 1.64]
        elif cols == 3:
            widths = [1.45, 2.70, 2.87] if self.number == 1 else [1.40, 3.10, 2.52]
        else:
            widths = [2.65, 4.37]
        tbl = self.doc.add_table(rows=1, cols=cols)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = False
        for col, width in zip(tbl.columns, widths):
            col.width = Inches(width)
        props = tbl._tbl.tblPr
        borders = OxmlElement("w:tblBorders")
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            b = OxmlElement(f"w:{edge}")
            b.set(qn("w:val"), "single")
            b.set(qn("w:sz"), "5")
            b.set(qn("w:color"), "D9D9D9")
            borders.append(b)
        props.append(borders)
        for i, row in enumerate(rows):
            cells = tbl.rows[0].cells if i == 0 else tbl.add_row().cells
            trpr = tbl.rows[i]._tr.get_or_add_trPr()
            trpr.append(OxmlElement("w:cantSplit"))
            if i == 0:
                trpr.append(OxmlElement("w:tblHeader"))
            for j, (cell, value) in enumerate(zip(cells, row)):
                cell.width = Inches(widths[j])
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                tcpr = cell._tc.get_or_add_tcPr()
                margins = OxmlElement("w:tcMar")
                for side in ("top", "bottom", "left", "right"):
                    edge = OxmlElement(f"w:{side}")
                    edge.set(qn("w:w"), "90" if side in ("top", "bottom") else "110")
                    edge.set(qn("w:type"), "dxa")
                    margins.append(edge)
                tcpr.append(margins)
                shade = OxmlElement("w:shd")
                shade.set(qn("w:fill"), "244B65" if i == 0 else "F2F6F8" if i % 2 == 0 else "FFFFFF")
                tcpr.append(shade)
                p = cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(1)
                p.paragraph_format.line_spacing = 1.22
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if cols == 4 and j else WD_ALIGN_PARAGRAPH.LEFT
                inline(p, value, size=9.8 if cols == 4 else 10, color="FFFFFF" if i == 0 else TEXT_COLOR)
                if i == 0:
                    for run in p.runs:
                        run.bold = True
        spacer = self.doc.add_paragraph()
        spacer.paragraph_format.space_after = Pt(3)
        spacer.paragraph_format.space_before = Pt(0)
        spacer.paragraph_format.line_spacing = Pt(3)
        if self.number == 1 and self.table_index == 2:
            self.image(FIG / "03-rmse-real-data.png",
                       "从已核验 CSV 直接绘制的实测结果。各面板量程和单位不同，误差线不是置信区间。")
        return tbl

    def code(self, lang, text):
        if self.number == 1 and text.startswith("保存来源快照"):
            self.paragraph("实际流程依次为保存来源快照、提取 PDF、导入同源审查记录、构建阅读包、严格验证，再用阅读包检索论文内容。")
            return
        if self.number == 2 and text.startswith("Pi GUI 输入框"):
            return
        if self.number == 2 and text.startswith("paper2agent/"):
            label = "完整 Skill 目录"
        elif text.startswith("/skill:paper2agent"):
            label = "可修改路径后使用的任务示例"
        else:
            label = {"powershell": "PowerShell", "python": "Python", "json": "验证结果", "text": "参数与依赖"}.get(lang, "代码")
        lp = self.doc.add_paragraph()
        lp.paragraph_format.space_before = Pt(6)
        lp.paragraph_format.space_after = Pt(4)
        lp.paragraph_format.keep_with_next = True
        set_font(lp.add_run(label), 9.5, True, "526579")
        code_lines = text.splitlines()
        for idx, value in enumerate(code_lines):
            p = self.doc.add_paragraph(style="Article Code")
            set_font(p.add_run(value or " "), 9.5, mono=True)
            p.paragraph_format.keep_together = True
            p.paragraph_format.keep_with_next = idx < len(code_lines) - 1
            if idx == len(code_lines) - 1:
                p.paragraph_format.space_after = Pt(10)
                if lang == "python" and text.startswith("rel = str(path.relative_to(root))"):
                    p.paragraph_format.keep_with_next = True

    def after_text(self, value):
        if self.number == 1:
            if value.startswith("2026 年 9 月 18 日至 19 日"):
                self.image(FIG / "01-tool-roles.png",
                           "两类工具的实际职责。本次科学实验由另外编写的基准脚本执行，不能归为任一工具自动完成的能力。")
            elif value.startswith("这也不能直接推导为"):
                self.image(FIG / "02-generation-vs-export.png",
                           "两轮 Paper2Any 测试使用了不同路径。9 月 19 日原生导出成功，不代表 9 月 18 日失败的模型生成链路已恢复。")
            elif value.startswith("因此，本轮结论是"):
                self.image(PREVIEW / "slide-07.png",
                           "Paper2Any 原生导出的实际页面。报告保留 FreeSolv 负结果；本次协议不同于原论文，图中差值只作描述性比较。")
        else:
            if value.startswith("本文以一次 Windows 实测"):
                self.image(FIG / "04-pi-call-architecture.png",
                           "Pi GUI 到本地转换脚本的调用架构示意。文件需通过被识别的路径打开；并非落盘后一定自动进入预览列表。")
            elif value.startswith("本次用到的环境不是"):
                self.image(FIG / "05-environments.png",
                           "本次实际环境的隔离方式。Paper2Skill 不依赖科学训练环境，Paper2Any 的 Python 环境也单独维护。",
                           width=6.8, space_before=0, caption_after=5)
            elif value.startswith("历史测试中，技能候选"):
                self.image(HISTORY / "gui-test/evidence/02-live-tools-and-task.png",
                           "2026 年 9 月 18 日真实 Pi GUI 截图。左侧任务在运行中可见，中间持续显示工具回执，输入框和停止按钮保留。右侧数值仅是当时界面快照。")
            elif value.startswith("为了排错，也为了能独立"):
                self.image(FIG / "06-review-flow.png",
                           "转换与审查流程示意。逐页审查不能由 review-aid 命令替代，严格验证通过也需要继续阅读具体状态与限制。")
            elif value.startswith("**长会话刷新失败。**"):
                self.image(HISTORY / "gui-test/evidence/06-markdown-preview.png",
                           "真实测试截图同时保留报告预览与右上角 contextBridge 错误。它说明界面可显示部分内容，但长会话全量状态刷新仍然失败。")

    def build(self, markdown_path, output_name):
        lines = Path(markdown_path).read_text(encoding="utf-8-sig").splitlines()
        i = 0
        while i < len(lines):
            line = lines[i]
            if not line.strip() or line == "---" or line.startswith("# ") or line == "文 / Stella":
                i += 1
                continue
            if line.startswith("##"):
                match = re.match(r"^(#{2,4}) (.*)$", line)
                heading = HEADINGS.get(match[2], match[2])
                heading = re.sub(r"[：:，,。！？!?“”「」/、]", " ", heading)
                p = self.doc.add_paragraph(heading, f"Heading {len(match[1]) - 1}")
                for run in p.runs:
                    set_font(run, color="000000")
                if len(match[1]) == 2:
                    self.chapter += 1
                i += 1
                continue
            if line.startswith("```"):
                lang = line[3:].strip()
                i += 1
                content = []
                while i < len(lines) and not lines[i].startswith("```"):
                    content.append(lines[i])
                    i += 1
                if i >= len(lines):
                    raise ValueError("Unclosed code fence")
                self.code(lang, "\n".join(content))
                i += 1
                continue
            if line.startswith("|"):
                table_lines = []
                while i < len(lines) and lines[i].startswith("|"):
                    table_lines.append(lines[i])
                    i += 1
                self.table(table_lines)
                continue
            if line.startswith("- "):
                self.paragraph(line[2:], "List Bullet")
                i += 1
                continue
            value = line.strip()
            if self.number == 2 and value == "调用关系可以简化为：":
                i += 1
                continue
            if value.startswith("资料核查日期："):
                value = "测试与资料核查截至 2026 年 9 月 19 日。本文所列命令与参数依据固定版本源码及保存的运行记录整理。跨平台安装与完整 GUI 回归不在本次验证范围内。"
            self.paragraph(value)
            self.after_text(value)
            i += 1
        destination = ROOT / output_name
        self.doc.save(destination)
        return {"path": str(destination), "figures": self.figures,
                "tables": self.table_index, "paragraphs": len(self.doc.paragraphs),
                "bytes": destination.stat().st_size}

if __name__ == "__main__":
    articles = [
        (1, "Paper2Agent 与 Paper2Any 实测", "论文阅读包 科学实验与演示文稿生成",
         "01-Paper2Agent与Paper2Any实测.md", "01-Paper2Agent与Paper2Any实测-图文版.docx"),
        (2, "Pi GUI 调用 Paper2Agent 技术实践", "Python 环境 技能安装与可核验的执行流程",
         "02-Pi-GUI调用Paper2Agent技术实践.md", "02-Pi-GUI调用Paper2Agent技术实践-图文版.docx"),
    ]
    outputs = [Article(number, title, subtitle).build(ROOT / markdown, output)
               for number, title, subtitle, markdown, output in articles]
    (BUILD / "document-build.json").write_text(json.dumps(outputs, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(outputs, ensure_ascii=False, indent=2))
