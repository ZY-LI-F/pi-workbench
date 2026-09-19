"""Verify Word structure, source fidelity, data cells and rendered page bounds."""
from __future__ import annotations
import csv
import hashlib
import json
import re
import xml.etree.ElementTree as ET
from pathlib import Path
from zipfile import ZipFile
from docx import Document
from docx.oxml.ns import qn
import pdfplumber

ROOT = Path(__file__).resolve().parents[1]
BUILD = ROOT / "word-build"
builds = json.loads((BUILD / "document-build.json").read_text(encoding="utf-8"))
reports = []

def expected_blocks(source):
    result = []
    for match in re.finditer(r"```[^\n]*\n(.*?)\n```", source, re.S):
        text = match[1]
        if text.startswith(("保存来源快照", "Pi GUI 输入框")):
            continue
        result.append(text)
    return result

for index, build in enumerate(builds, start=1):
    path = Path(build["path"])
    source_path = path.with_name(path.name.replace("-图文版.docx", ".md"))
    source = source_path.read_text(encoding="utf-8-sig")
    doc = Document(path)
    assert len(doc.inline_shapes) == (4 if index == 1 else 5)
    assert len(doc.tables) == (3 if index == 1 else 2)
    assert doc.paragraphs[0].style.name == "Title"
    for p in doc.paragraphs:
        if p.style.name in ("Title", "Subtitle") or p.style.name.startswith("Heading"):
            assert all(r.font.color.rgb and str(r.font.color.rgb) == "000000" for r in p.runs), p.text
            assert re.fullmatch(r"[\w\s]+", p.text), p.text
            assert not p._p.xpath("./w:pPr/w:pBdr")
    code_groups = []
    group = []
    for p in doc.paragraphs:
        if p.style.name == "Article Code":
            group.append(p.text if p.text.strip() else "")
        elif group:
            code_groups.append("\n".join(group))
            group = []
    if group:
        code_groups.append("\n".join(group))
    assert code_groups == expected_blocks(source), "A code block changed in conversion"
    md_tables = []
    for block in re.findall(r"(?m)^\|[^\n]+(?:\n\|[^\n]+)+", source):
        rows = [[c.strip() for c in line.strip().strip("|").split("|")] for line in block.splitlines()]
        rows = [row for row in rows if not all(re.fullmatch(r":?-+:?", c) for c in row)]
        md_tables.append(rows)
    assert len(md_tables) == len(doc.tables)
    for table, expected in zip(doc.tables, md_tables):
        expected = [[re.sub(r"`([^`]+)`", r"\1", cell) for cell in row] for row in expected]
        actual = [[cell.text for cell in row.cells] for row in table.rows]
        assert actual == expected, (actual, expected)
        assert table.rows[0]._tr.xpath("./w:trPr/w:tblHeader")
        assert all(row._tr.xpath("./w:trPr/w:cantSplit") for row in table.rows)
    captions = [p.text for p in doc.paragraphs if p.style.name == "Caption"]
    assert len(captions) == len(doc.inline_shapes)
    assert all(c.startswith(f"图 {n}  ") for n, c in enumerate(captions, 1))
    with ZipFile(path) as package:
        assert package.testzip() is None
        xml = package.read("word/document.xml").decode()
        assert "\ufffd" not in xml
        assert not re.search(r"TODO|TBD|PLACEHOLDER|turn\d+(?:search|view)", xml)
        image_hashes = {hashlib.sha256(package.read(n)).hexdigest()
                        for n in package.namelist() if n.startswith("word/media/")}
        for figure in build["figures"]:
            assert hashlib.sha256(Path(figure["path"]).read_bytes()).hexdigest() in image_hashes
        for shape in doc.inline_shapes:
            assert shape._inline.docPr.get("descr")
            assert shape._inline.docPr.get("title")
        rels = ET.fromstring(package.read("word/_rels/document.xml.rels"))
        links = {r.attrib["Target"] for r in rels
                 if r.attrib.get("Type", "").endswith("/hyperlink")}
        source_links = set(re.findall(r"\[[^\]]+\]\((https?://[^)]+)\)", source))
        assert links == source_links
        assert not any(r.attrib.get("TargetMode") == "External" and r.attrib["Type"].endswith("/image") for r in rels)
    render = BUILD / f"render-{index:02d}-v3"
    pdf = render / (path.stem + ".pdf")
    page_report = []
    with pdfplumber.open(pdf) as rendered_pdf:
        pages = rendered_pdf.pages
        assert len(pages) == len(list(render.glob("page-*.png")))
        for n, page in enumerate(pages, 1):
            words = page.extract_words()
            assert words, f"Blank page {n}"
            outliers = [w["text"] for w in words if
                        w["x0"] < 40 or w["x1"] > page.width - 35 or
                        w["top"] < 25 or w["bottom"] > page.height - 8]
            assert not outliers, (index, n, outliers)
            page_report.append({"page": n, "words": len(words), "text_bounds_inside_page": True})
    reports.append({
        "path": str(path), "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
        "bytes": path.stat().st_size, "pages": len(pages),
        "figures": len(doc.inline_shapes), "native_tables": len(doc.tables),
        "native_code_blocks": len(code_groups), "source_hyperlinks": len(links),
        "original_image_bytes_preserved": True, "source_code_and_tables_preserved": True,
        "page_checks": page_report,
    })

rows = list(csv.DictReader((ROOT.parent / "moleculenet-reproduction-20260918/verification/summary.csv").open(encoding="utf-8-sig")))
lookup = {(r["dataset"], r["model"]): r for r in rows if r["split"] == "random"}
doc = Document(builds[0]["path"])
result_table = doc.tables[1]
for i, dataset in enumerate(("esol", "freesolv", "lipophilicity"), 1):
    for j, model in enumerate(("rf", "krr", "gc"), 1):
        row = lookup[(dataset, model)]
        expected = f'{float(row["test_rmse_mean"]):.3f} ± {float(row["test_rmse_sample_sd"]):.3f}'
        assert result_table.cell(i, j).text == expected
result = {
    "status": "passed", "documents": reports,
    "rmse_cells_matched_to_csv": 9,
    "rendering": "Isolated signed LibreOffice 26.2.6 administrative extraction and bundled Poppler via documents render_docx.py",
    "visual_review": "Recorded separately after inspecting every final page PNG",
    "not_claimed": ["Microsoft Word desktop testing", "WPS desktop testing", "new scientific or model test execution"]
}
(BUILD / "word-validation.json").write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
print(json.dumps({k:v for k,v in result.items() if k != "documents"}, ensure_ascii=False, indent=2))
for report in reports:
    print(json.dumps({k:v for k,v in report.items() if k != "page_checks"}, ensure_ascii=False, indent=2))
